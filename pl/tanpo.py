#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""店舗ごとの単体PL（店長さんが見るスプレッドシート）へ、人件費を写す

    GOOGLE_SA_KEY_FILE=/tmp/sa.json python3 tanpo.py --dry-run   # 見るだけ
    GOOGLE_SA_KEY_FILE=/tmp/sa.json python3 tanpo.py             # 書く

--- 利用者指示 2026-10-18 ---------------------------------------------------
「もも焼きJAPANの単体のPL、ここに店長の人件費、福利厚生費を入れといて」
→「はい、毎月お願い」「他の店舗、焼きたて屋、ハナ、タコハイ、りゅうちゃんにも入れて」

元は 22期の全体PL（fill22.py が作る 損益計算書_22期.xlsx）の各店タブ。
写すのは「人件費（店長）」と「法定福利費」だけ（単体PLに『福利厚生費』の行は無いので法定福利費）。
タコハイの単体PLは店長の行が「人件費（社員）」という名前。

★書き込みの決まり（振込一覧と同じ）
  ・空のセルにだけ書く。同じ値が入っていれば何もしない。
  ・違う値・数式が入っていたら書かずに「食い違い」として出す（利用者に聞く）。
  ・全体PLのほうが空（0）の月は何もしない（給料一覧表がまだ等）。
行は B列の見出し、月は 4行目の見出しで探す（行番号は店ごとに違う）。
"""
import os
import sys

import openpyxl

XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "損益計算書_22期.xlsx")

# 全体PLのタブ: (単体PLのファイル, タブ, {全体PLの行: 単体PLの行})
STORES = {
    "もも焼きJAPAN": ("1v44EK340mt-BY2z-w6M12dfzvYaegUXR6Wx9FqtXThA", "2026PL",
                   {"人件費（店長）": "人件費（店長）", "法定福利費": "法定福利費"}),
    "りゅうちゃん": ("1tivc_3ijpZNuMu0NEfdn699hxSZJSR_U9HvUgTpE3vE", "2026PL",
                {"人件費（店長）": "人件費（店長）", "法定福利費": "法定福利費"}),
    "焼きたて屋": ("14jjx9OM_YqPTna9e75FZCredSY3tfH4-tgyo7vXmlYk", "2026PL",
               {"人件費（店長）": "人件費（店長）", "法定福利費": "法定福利費"}),
    "タコとハイボール": ("1Kbz8KD8d5ezMgrbCDMQhk05H1ub58CjwJdc-vn61ECQ", "PL2026",
                  {"人件費（店長）": "人件費（社員）", "法定福利費": "法定福利費"}),
    "韓国酒場ハナ": ("1t5lBhKl-BOpbO487OpYu7xZtv_gkYwLd--GGXYOsHGA", "2026PL",
                {"人件費（店長）": "人件費（店長）", "法定福利費": "法定福利費"}),
}
# 利用者の答え待ちで写さない店（2026-10-18）
#   焼きたて屋の単体PLは 9月に 店長427,159・アルバイト40,620 が手で入っている。
#   全体PLは アルバイト427,159・法定福利費40,620（店長は空）。同じ2つの数字が1行ずつずれているので、
#   法定福利費に40,620を足すと二重になる。どちらに合わせるか聞いてから外す。
HOLD = {"焼きたて屋"}
MONTHS = ["9月", "10月", "11月", "12月", "1月", "2月", "3月", "4月", "5月", "6月", "7月", "8月"]


def col(i0):
    s = ""
    i0 += 1
    while i0:
        i0, r = divmod(i0 - 1, 26)
        s = chr(65 + r) + s
    return s


def source():
    """全体PL → {(タブ, 行, 月): 値}"""
    wb = openpyxl.load_workbook(XLSX)
    out = {}
    for tab, (_, _, rows) in STORES.items():
        ws = wb[tab]
        for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
            if r[0].value in rows:
                for i, m in enumerate(MONTHS):
                    v = r[2 + i].value
                    if isinstance(v, (int, float)) and v:
                        out[(tab, r[0].value, m)] = int(v)
    return out


def num(s):
    s = str(s).replace("¥", "").replace(",", "").replace("▲", "-").strip()
    try:
        return int(float(s))
    except ValueError:
        return None


def main():
    dry = "--dry-run" in sys.argv
    from googleapiclient.discovery import build
    from google.oauth2.service_account import Credentials
    creds = Credentials.from_service_account_file(
        os.environ["GOOGLE_SA_KEY_FILE"], scopes=["https://www.googleapis.com/auth/spreadsheets"])
    api = build("sheets", "v4", credentials=creds, cache_discovery=False).spreadsheets()

    src = source()
    writes, same, clash = [], 0, []
    for tab, (sid, sheet, rows) in STORES.items():
        if tab in HOLD:
            print(f"  （{tab} は保留中なので写さない）")
            continue
        rng = f"'{sheet}'!A1:Z80"
        frm = api.values().get(spreadsheetId=sid, range=rng, valueRenderOption="FORMULA").execute().get("values", [])
        shown = api.values().get(spreadsheetId=sid, range=rng).execute().get("values", [])
        head = frm[3] if len(frm) > 3 else []
        mcol = {h: i for i, h in enumerate(head) if h in MONTHS}
        rix = {r[1]: i for i, r in enumerate(frm) if len(r) > 1 and r[1] in rows.values()}
        for prow, srow in rows.items():
            if srow not in rix:
                clash.append(f"{tab}: 単体PLに『{srow}』の行が無い")
                continue
            for m in MONTHS:
                v = src.get((tab, prow, m))
                if v is None:
                    continue
                ri, ci = rix[srow], mcol[m]
                cur = frm[ri][ci] if ci < len(frm[ri]) else ""
                curv = shown[ri][ci] if ri < len(shown) and ci < len(shown[ri]) else ""
                cell = f"{col(ci)}{ri + 1}"
                if cur in ("", None):
                    writes.append((tab, sid, sheet, cell, srow, m, v))
                elif num(curv) == v:
                    same += 1
                else:
                    clash.append(f"{tab} {srow} {m}（{cell}）: 単体PLに {curv} が入っている／全体PLは {v:,}")

    for tab, _, _, cell, srow, m, v in writes:
        print(f"  書く  {tab} {srow} {m}（{cell}）{v:,}")
    for c in clash:
        print(f"  ★食い違い {c}")
    print(f"書く {len(writes)} ／ 同じ値で入り済み {same} ／ 食い違い {len(clash)}"
          + ("（dry-run: 書き込みなし）" if dry else ""))
    if dry or not writes:
        return
    by = {}
    for tab, sid, sheet, cell, _, _, v in writes:
        by.setdefault(sid, []).append({"range": f"'{sheet}'!{cell}", "values": [[v]]})
    for sid, data in by.items():
        api.values().batchUpdate(spreadsheetId=sid,
                                 body={"valueInputOption": "USER_ENTERED", "data": data}).execute()
    # 読み直し
    for tab, sid, sheet, cell, srow, m, v in writes:
        got = api.values().get(spreadsheetId=sid, range=f"'{sheet}'!{cell}").execute().get("values")
        assert got and num(got[0][0]) == v, (tab, cell, got)
    print("読み直して全件一致")


if __name__ == "__main__":
    main()
