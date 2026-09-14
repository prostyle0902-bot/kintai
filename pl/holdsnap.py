#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""保留リストのスナップショットを取り、前回から増えた保留を検出する

    python3 holdsnap.py            # いまの保留を表示するだけ
    python3 holdsnap.py --check    # 前回スナップショットと比べる（増えていたら exit 1）
    python3 holdsnap.py --save     # いまの保留をスナップショットとして記録する

なぜ要るか
----------
定期チェック（Routine）でPLまで自動で反映するようになった（2026-09-14）。
利用者の選択は「**保留が増えていなければ Drive まで自動で反映、増えたら止めて報告**」。

pl/README.md の大原則「判定できなかったものは必ず保留リストへ出す。推測で埋めない」
「転記前に一覧を利用者に見せて確認を取る」を、この仕組みで守る:

  - 保留は常に90件前後ある（8月の未着資料・カード一括など）。だから
    「保留ゼロなら」ではなく「**前回から増えていなければ**」で判定する。
  - 増えた保留＝利用者がまだ見ていない判断事項。これがあるうちは push しない。

`hold_snapshot.tsv` は「利用者に見せて Drive へ反映済み」の状態を表す。
push が通ったあとに `--save` で更新する。リポジトリにコミットしておくこと
（コンテナはセッションごとに消えるので、ファイルで持たないと比較できない）。
"""
import collections
import os
import sys

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "損益計算書_21期テスト版.xlsx")
SNAP = os.path.join(BASE, "hold_snapshot.tsv")
COLS = 6          # 区分 / 日付・月 / 店舗 / 取引先 / 金額(税込) / 理由・メモ


def current(path=XLSX):
    """保留リストシートを {1行=タブ区切り文字列: 件数} の多重集合で返す。

    並び順は fill2.py の出力順に左右されるので、比較は集合でやる。
    ★同じ文面の行が複数あることがある（定額の理由文など）ので、件数まで持つ。
      集合にすると重複ぶんの増減を取りこぼす。
    """
    if not os.path.exists(path):
        sys.exit(f"{path} が無い。先に python3 fill2.py を流すこと")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "保留リスト" not in wb.sheetnames:
        sys.exit("保留リストシートが無い。fill2.py の出力が変わっていないか見ること")
    hs = wb["保留リスト"]
    rows = collections.Counter()
    for r in hs.iter_rows(min_row=3, max_col=COLS, values_only=True):
        cells = ["" if v is None else str(v).strip() for v in r]
        if not any(cells):
            continue
        # 「保留はありません」の1行だけのときは A列にしか値が無い
        if cells[0].startswith("保留はありません"):
            continue
        rows["\t".join(cells)] += 1
    wb.close()
    return rows


def load_snapshot():
    if not os.path.exists(SNAP):
        return None
    with open(SNAP, encoding="utf-8") as f:
        return collections.Counter(ln.rstrip("\n") for ln in f if ln.strip())


def save_snapshot(rows):
    with open(SNAP, "w", encoding="utf-8") as f:
        for ln in sorted(rows.elements()):
            f.write(ln + "\n")


def show(rows, label):
    print(f"--- {label}（{sum(rows.values())}件）")
    for ln in sorted(rows.elements()):
        print("   ", ln.replace("\t", " / "))


def main():
    args = sys.argv[1:]
    rows = current()

    if "--save" in args:
        save_snapshot(rows)
        print(f"保留 {sum(rows.values())} 件を {os.path.basename(SNAP)} に記録した")
        return 0

    if "--check" not in args:
        show(rows, "いまの保留")
        return 0

    old = load_snapshot()
    if old is None:
        show(rows, "いまの保留")
        print(f"\n★{os.path.basename(SNAP)} が無い（初回）。"
              "中身を利用者に見せて確認を取ってから --save すること")
        return 1

    added = rows - old
    gone = old - rows
    if gone:
        show(gone, "解消した保留")
    if added:
        show(added, "★増えた保留（利用者に報告し、判断をもらうまで push しない）")
        return 1
    print(f"保留は増えていない（{sum(rows.values())}件、"
          f"うち解消 {sum(gone.values())}件）。push してよい")
    return 0


if __name__ == "__main__":
    sys.exit(main())
