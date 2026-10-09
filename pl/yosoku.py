#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""予測表（Googleスプレッドシート「YYYY年M月_予測表」）の飲食事業部に、店舗ごとの日別売上（税抜）を入れる

--- 利用者指示 2026-10-09 ---------------------------------------------------
「予測表の飲食事業部の各店舗に売上入れてるんだけど、各店舗CSVを見て売上入力してほしい」
→ 9月を入れた →「毎月、お願いします！」

    GOOGLE_SA_KEY_FILE=/tmp/sa.json python3 yosoku.py 2026-10            # 入れる
    GOOGLE_SA_KEY_FILE=/tmp/sa.json python3 yosoku.py 2026-10 --dry-run  # 見るだけ

--- どこに入れるか ------------------------------------------------------------
シート「ＭＱ粗利予定表」の G〜J 列（飲食事業部）。1日＝10行のかたまりで、
G列にその日の日付（MM/DD）、I列に店舗名、J列「売上（税抜）」。
行の位置は毎回 G列・I列を読んで決める（月によって31日目のかたまりがある・行がずれることがあるため）。
★横丁の行（豚骨流星群・はな など）は元データが無いので入れない。

--- 元データ（PLと同じもの） -------------------------------------------------
  エアレジ日別CSV  uriage/<YYMM月>/<店>_<YYYYMMDD>-<YYYYMMDD>.csv
      税抜 ＝ 10%標準の税込を逆算 ＋ 8%軽減の税込を逆算（airregi.back。PLの消費税と同じ計算）
  焼きたて屋       uriage/<YYMM月>/焼きたて屋_税抜.xlsx（かめやのFC月間売上集計一覧表・税抜版）の日別「ＦＣ合計」
各店の月合計は PL の売上（税抜）と1円まで合うはず（合わなければ止まる）。

--- 決まり --------------------------------------------------------------------
・空のセルにだけ書く。すでに同じ値ならそのまま。違う値が入っていたら書かずに報告する（利用者の手入力を消さない）。
・売上が無い日（休み）は空欄のまま。
・元データがそろっていない店は飛ばす（届いてから次の回で入る）。
"""
import csv
import glob
import io
import os
import sys
import warnings

BASE = os.path.dirname(os.path.abspath(__file__))
SHEET = "ＭＱ粗利予定表"
# CSVの店名 → 予測表の店名
CSV_STORES = {"もも焼き": "もも焼きJAPAN", "りゅうちゃん": "りゅうちゃん", "タコハイ": "タコハイ",
              "ハナ": "ハナ", "十三里屋": "十三里屋"}
YAKITATE = "焼きたて屋"


def daily(ym):
    """{(予測表の店名, 'MM/DD'): 税抜}"""
    import airregi
    import openpyxl
    y, m = ym.split("-")
    folder = os.path.join(BASE, "uriage", f"{y[2:]}{m}月")
    out, totals = {}, {}
    for path in sorted(glob.glob(os.path.join(folder, f"*_{y}{m}01-*.csv"))):
        name = os.path.basename(path).split("_")[0]
        if name not in CSV_STORES:
            continue
        st = CSV_STORES[name]
        rows = list(csv.DictReader(io.open(path, encoding="cp932").read().splitlines()))
        tot = 0
        for r in rows:
            d = r["集計期間"]
            assert d[:6] == f"{y}{m}", f"{path}: {d} が {ym} でない"
            v10, v8 = int(r["売上（10%標準）"]), int(r["売上（8%軽減）"])
            assert v10 + v8 == int(r["売上"]), f"{path} {d}: 税率別の合計が売上と違う"
            ex = airregi.back(v10, 10) + airregi.back(v8, 8)
            if ex:
                out[(st, f"{d[4:6]}/{d[6:8]}")] = ex
                tot += ex
        inc, tax, *_ = airregi._csv(path)
        assert tot == inc - tax, f"{st}: 日別の税抜 {tot:,} が月の税抜 {inc - tax:,} と合わない"
        totals[st] = tot
    hit = glob.glob(os.path.join(folder, "焼きたて屋_*税抜*.xlsx"))
    if hit:
        warnings.filterwarnings("ignore")
        ws = openpyxl.load_workbook(hit[0], data_only=True)["ＦＣ店"]
        title = str(ws.cell(1, 1).value or "")
        assert f"{y}年{m}月" in title and "税抜" in title, f"{hit[0]}: 表題が {title}"
        tot, total_row = 0, None
        for r in ws.iter_rows(min_row=4, values_only=True):
            if isinstance(r[0], int) and r[4] not in (None, ""):
                v = int(str(r[4]).replace(",", ""))
                if v:
                    out[(YAKITATE, f"{m}/{r[0]:02d}")] = v
                    tot += v
            if r[0] and str(r[0]).replace("　", "") == "合計":
                total_row = int(str(r[4]).replace(",", ""))
        assert total_row == tot, f"焼きたて屋: 日別の合計 {tot:,} が表の合計 {total_row} と合わない"
        totals[YAKITATE] = tot
    return out, totals


def find_sheet(drv, ym):
    y, m = ym.split("-")
    name = f"{y}年{int(m)}月_予測表"
    got = drv.files().list(q=f"name = '{name}' and trashed = false and "
                             "mimeType = 'application/vnd.google-apps.spreadsheet'",
                           fields="files(id,name)", supportsAllDrives=True,
                           includeItemsFromAllDrives=True).execute()["files"]
    assert len(got) == 1, f"「{name}」が {len(got)} 件見つかった（1件のはず）"
    return got[0]["id"], name


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 1:
        sys.exit("使い方: python3 yosoku.py YYYY-MM [--dry-run]")
    ym, dry = args[0], "--dry-run" in sys.argv
    vals, totals = daily(ym)
    if not vals:
        print(f"{ym}: 元データがまだ無い")
        return
    import push
    from googleapiclient.discovery import build
    cr = push._creds()
    drv = build("drive", "v3", credentials=cr)
    svc = build("sheets", "v4", credentials=cr)
    sid, name = find_sheet(drv, ym)
    T = f"'{SHEET}'"
    a1 = svc.spreadsheets().values().get(spreadsheetId=sid, range=f"{T}!A1").execute()["values"][0][0]
    assert str(a1) == str(int(ym[5:])), f"{name}: A1 の月が {a1}"
    grid = svc.spreadsheets().values().get(spreadsheetId=sid, range=f"{T}!G1:J320",
                                           valueRenderOption="UNFORMATTED_VALUE").execute()["values"]
    shown = svc.spreadsheets().values().get(spreadsheetId=sid, range=f"{T}!G1:G320").execute()["values"]
    pos, cur_day = {}, None
    for i, r in enumerate(grid, start=1):
        r = list(r) + [""] * 4
        g = shown[i - 1][0].strip() if i - 1 < len(shown) and shown[i - 1] else ""
        if i >= 4 and g:
            cur_day = g
        if i >= 4 and cur_day and str(r[2]).strip():
            pos[(str(r[2]).strip(), cur_day)] = (i, r[3])
    missing = [k for k in vals if k not in pos]
    assert not missing, f"予測表に行が見つからない: {missing[:5]}"
    write, same, clash = [], 0, []
    for k, v in sorted(vals.items(), key=lambda kv: pos[kv[0]][0]):
        row, now = pos[k]
        if now in ("", None):
            write.append((row, v))
        elif now == v:
            same += 1
        else:
            clash.append((k, row, now, v))
    print(f"{name}: 書く {len(write)} ／ 同じ値で入り済み {same} ／ 違う値が入っている {len(clash)}")
    for st in [YAKITATE] + list(CSV_STORES.values()):
        if st in totals:
            print(f"  {st:<12} 月計（税抜） {totals[st]:>10,}")
    for k, row, now, v in clash:
        print(f"  ★書かなかった: {k[0]} {k[1]}（J{row}）いま {now} ／ CSV {v:,}")
    if dry or not write:
        return
    svc.spreadsheets().values().batchUpdate(spreadsheetId=sid, body={
        "valueInputOption": "RAW",
        "data": [{"range": f"{T}!J{row}", "values": [[v]]} for row, v in write]}).execute()
    back = svc.spreadsheets().values().get(spreadsheetId=sid, range=f"{T}!J1:J320",
                                           valueRenderOption="UNFORMATTED_VALUE").execute()["values"]
    bad = [(row, v) for row, v in write if not (row - 1 < len(back) and back[row - 1]
                                                 and back[row - 1][0] == v)]
    assert not bad, f"書いたのに読み直すと違う: {bad[:5]}"
    print(f"  書き込みを読み直して確認した（{len(write)}セル）")


if __name__ == "__main__":
    main()
