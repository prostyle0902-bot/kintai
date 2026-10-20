#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""棚卸し（期首・期末）

--- どこから取ったか -----------------------------------------------------
Googleドライブ。2種類ある。

① ※Prostyle株式会社/21期_棚卸し表 .xlsx   id=17NZiAKPDma17Ayj6uUOJxLoMk8Bksts7
   8タブ（ビルメン・鳩／タコハイ／焼きたて屋／もも焼き／リュウちゃん／
   神栖横丁／十三里屋／ハナ）。2025年9月に入力した実地棚卸＝【21期の期首】。
   ★韓国酒場ハナの月次棚卸表に「20250831」というシートがあって 144,577。
     ①のハナと1円まで一致した。つまり①の基準日は【2025年8月31日】。
   ★①の合計行は十三里屋だけ ¥180 と壊れている（タブ自身の合計 273,232 が正）。

② 各店フォルダ（共有）の月次棚卸表。シート名＝YYYYMMDD、
   O3に棚卸年月日、R3に合計（=J57+T57 の数式）。
       2.焼きたて屋/棚卸表_焼きたて屋.xlsx              id=1fJPUbw3WNxtBtEiB8jOlYwqCsdO3-N75
       3.タコとハイボール/棚卸表_タコとハイボール.xlsx    id=1y2Gc-c3PoLSkqngJT2bZVqWKKWF8UBdI
       4.もも焼JAPAN/棚卸表_もも焼JAPAN.xlsx          id=1GHGnIXPrzt_KXVS9B4uDn5C0tFkulFII
       5.大衆酒場りゅうちゃん/棚卸表_大衆酒場りゅうちゃん.xlsx id=1B4QaxjQ2YUPmfBc9bOn9Jwx3NhrrqQ-J
       6.韓国酒場ハナ/棚卸表_韓国酒場ハナ.xlsx           id=16M2TxUwFBvyH-Pye9En34esBs4bUmLko

--- ★端数は Excel と同じ四捨五入（ROUND_HALF_UP）------------------------
R3は小数を持っている（例 251383.783）。シート上の表示は 251,384。
既存21期PLもその表示値を使っていた（タコハイ 200180.5 → 200,181）。
Python の round() は銀行丸めで 200,180 になるので使っていない。
ここに書いてあるのは四捨五入した後の整数。

--- ★別紙シートを足してはいけない -------------------------------------
りゅうちゃんともも焼きは「ドリンク」「その他日本酒」が別シートになっているが、
本紙の7行目あたりに「ドリンク小計（別紙）」という行があって、その合計が
そのまま入っている。つまり本紙のR3＝本紙＋別紙。足すと二重になる（実測で確認）。
    りゅうちゃん 20260131 本紙367,970 のうち ドリンク小計（別紙）225,823
    もも焼き   20251231 本紙484,880 のうち ドリンク1 306,916 ＋ドリンク2 37,722

--- どう並べるか ---------------------------------------------------------
    期首[9月]  = ①の実地棚卸（2025年8月31日）
    期末[m]    = ②の m月末の実地棚卸。無い月は前月を持ち越す
    期首[m]    = 期末[m-1]
既存21期PLは「最初に数えた月の値」を9月まで遡って置いていただけだった。
たとえば焼きたて屋は2026年2月に初めて数えた295,260を9月〜2月に置いていたが、
本当の期首は212,141。②の実地がある月は既存PLとほぼ一致していたので動かない。

--- ★8月（21期の期末）を入れた（2026-09-10）----------------------------
9/1〜9/4 に各店の20260831シートが埋まったので取り込んだ。
    タコとハイボール  161,096（R3 161095.588…）
    焼きたて屋       267,392（R3 267392.392…）
    りゅうちゃん     387,537（R3 387536.6）
    もも焼きJAPAN    227,901（R3 227901）
    さわら十三里屋   347,962 ★十三里屋は21期で【初めての実地棚卸】
        （新しいスプレッドシート「棚卸表_さわら十三里屋_20260831」
          id=1XZ5gGE2Us0DlziyXl6Nq1yHgFmiy6qIpgk8X4pO8oHU。
          小計 257,059＋90,903＝347,962 で合計欄と一致）
7月の値も同じやり方で読み直して GETSUJI と1円まで一致することを確かめた
（もも焼き242,473／タコハイ162,999／りゅうちゃん408,016／焼きたて屋251,384）。

--- ★韓国酒場ハナの8月も入った（2026-09-11）-----------------------------
棚卸表_韓国酒場ハナ.xlsx（id=16M2TxUwFBvyH-Pye9En34esBs4bUmLko）は
9/10 時点ではサービスアカウントに共有されておらず404で読めなかったが、
利用者が共有してくれたので読めるようになった。
    韓国酒場ハナ 20260831 シート R3=182459.9 → 182,460（J57 95,560.4＋T57 86,899.5）
既存の月（2月67,248／3月78,265／4月100,821／5月152,225／6月147,932／7月142,618）と
期首の20250831（144,577）も同じやり方で読み直して、GETSUJI・KISHU と1円まで一致した。
ハナに別紙シートは無い（シートは日付8枚だけ）。

--- もも焼きの9月・10月・11月（利用者確認 2026-09-01）------------------
既存21期PLに438,111／463,688／484,615が入っている。毎月ちがう数字なので
実際に数えた値とみられるが、②の棚卸表は20251231からしかシートが無く裏が取れない。
利用者確認「もも焼きの9・10・11月はナシ。既存のままでOK」→ 既存PLのまま。
期首12月＝期末11月＝484,615、期首1月＝期末12月＝484,880 でつながっている。

--- 月次の棚卸表が無い部署 ----------------------------------------------
十三里屋・神栖横丁・業務課・鳥害対策課 … 月次がまったく無い。
    ①の値を9月〜7月へそのまま流している。期首＝期末なので売上原価には
    影響しない（「分からない月は動かさない」という置き方）。

--- ★ビルメン・鳩タブの割り振り（利用者確認 2026-09-01）----------------
①の「ビルメン・鳩」タブ 1,813,683 は業務課と鳥害対策課がまざっている。
品目を数えて内訳を出した（合計は1円まで一致）:
    ビルメン薬剤・資材   1,080,271 … つやげん、光触媒、フェイスガードほか
    事務用品               15,828 … コピー用紙、ペーパータオルほか
    鳩対策の資材          345,644 … ピーコンネット、スパイク、パイラックほか
    ネットカーテンセット   371,940 … 利用者確認「ネットカーテンは鳩対策」
  → 業務課    1,080,271 ＋ 15,828 ＝ 1,096,099
     鳥害対策課  345,644 ＋ 371,940 ＝   717,584
"""

import os

MONTHS = ["9月", "10月", "11月", "12月", "1月", "2月", "3月", "4月", "5月", "6月", "7月", "8月"]

# ① 21期_棚卸し表（2025年8月31日の実地棚卸）＝ 21期の期首
KISHU = {
    "タコとハイボール": 228550,
    "焼きたて屋":       212141,
    "りゅうちゃん":     287557,
    "もも焼きJAPAN":   393397,
    "韓国酒場ハナ":     144577,
    "神栖横丁":          50357,
    "さわら十三里屋":   273232,
    "業務課":          1096099,   # ビルメン薬剤・資材1,080,271＋事務用品15,828
    "鳥害対策課":       717584,   # 鳩対策の資材345,644＋ネットカーテン371,940
}

# ② 月次棚卸表。{タブ: {PL列: その月末の実地棚卸}}
GETSUJI = {
    "タコとハイボール": {"12月": 200181, "1月": 117188, "2月": 221295, "3月": 207165,
                   "4月": 158124, "5月": 167066, "6月": 141000, "7月": 162999,
                   "8月": 161096},
    "焼きたて屋":     {"2月": 295260, "3月": 230231, "4月": 248524,
                   "5月": 217061, "6月": 201814, "7月": 251384, "8月": 267392},
    "りゅうちゃん":   {"1月": 367970, "2月": 347359, "3月": 443977, "4月": 402606,
                   "5月": 396109, "6月": 361102, "7月": 408016, "8月": 387537},
    "もも焼きJAPAN": {"12月": 484880, "1月": 400164, "2月": 301988, "3月": 263531,
                   "4月": 257611, "5月": 240167, "6月": 214183, "7月": 242473,
                   "8月": 227901},
    "韓国酒場ハナ":   {"2月": 67248, "3月": 78265, "4月": 100821,
                   "5月": 152225, "6月": 147932, "7月": 142618, "8月": 182460},
    # ★十三里屋は21期で初めての実地棚卸（2026年8月31日）。9月〜7月は①の持ち越し
    "さわら十三里屋": {"8月": 347962},
}

# ★8月を入れないタブ。棚卸表が読めない店があればここに入れる。
#   期首だけ入れると 売上原価(a)＝仕入合計＋期首−期末 が棚卸ぶんまるごと多く出るため。
#   韓国酒場ハナは 2026-09-11 に棚卸表がサービスアカウントへ共有されたので外した。
NO_AUG = set()

# ★既存21期PLの期末をそのまま残す月。
#   もも焼きの9月・10月・11月（438,111／463,688／484,615）は毎月ちがう数字で、
#   持ち越しではなく実際に数えた値。ドライブの棚卸表は20251231からしか
#   シートが無いので出所は追えないが、利用者確認 2026-09-01
#   「もも焼きの9・10・11月はナシ。既存のままでOK」で既存PLのままに確定した。
KEEP_EXIST = {"もも焼きJAPAN": {"9月", "10月", "11月"}}

SRC1 = "Googleドライブ ※Prostyle株式会社/21期_棚卸し表 .xlsx（2025年8月31日の実地棚卸）"
SRC2 = {
    "タコとハイボール": "Googleドライブ 3.タコとハイボール/棚卸表_タコとハイボール.xlsx",
    "焼きたて屋":       "Googleドライブ 2.焼きたて屋/棚卸表_焼きたて屋.xlsx",
    "りゅうちゃん":     "Googleドライブ 5.大衆酒場りゅうちゃん/棚卸表_大衆酒場りゅうちゃん.xlsx",
    "もも焼きJAPAN":   "Googleドライブ 4.もも焼JAPAN/棚卸表_もも焼JAPAN.xlsx",
    "韓国酒場ハナ":     "Googleドライブ 6.韓国酒場ハナ/棚卸表_韓国酒場ハナ.xlsx",
    "さわら十三里屋":   "Googleドライブ 棚卸表_さわら十三里屋_20260831（8月のみ。9月〜7月は①の持ち越し）",
}
_BESSHI = {"りゅうちゃん": "★本紙のR3にはドリンク（別紙）の小計が行として入っているので、"
                      "ドリンクシートを足すと二重になる",
           "もも焼きJAPAN": "★本紙のR3にはドリンク・その他日本酒（別紙）の小計が行として"
                        "入っているので、別紙を足すと二重になる"}


def rows():
    """(タブ, PL行, 月, 金額, 元ファイル, メモ) を列挙。"""
    for tab, kishu in KISHU.items():
        getsuji = GETSUJI.get(tab, {})
        keep = KEEP_EXIST.get(tab, set())
        src2 = SRC2.get(tab, SRC1)
        besshi = _BESSHI.get(tab, "")
        months = MONTHS[:-1] if tab in NO_AUG else MONTHS
        # 期末を月順に組む（実地がある月はその値、無い月は前月を持ち越し）
        last, kimatsu = kishu, {}
        for m in months:
            if m in getsuji:
                last = getsuji[m]
            kimatsu[m] = last
        for i, m in enumerate(months):
            # --- 期首 ---
            if i == 0:
                yield (tab, "期首棚卸し", m, kishu, SRC1,
                       f"2025年8月31日の実地棚卸 {kishu:,}（21期の期首）。"
                       "★既存21期PLは『最初に数えた月の値』を9月まで遡って置いていたので、"
                       "そこを実地の数字に直している")
            elif MONTHS[i - 1] not in keep:
                pm = MONTHS[i - 1]
                pk = "実地" if pm in getsuji else "持ち越し"
                yield (tab, "期首棚卸し", m, kimatsu[pm], src2,
                       f"前月（{pm}）末の棚卸 {kimatsu[pm]:,}（{pk}）をそのまま期首に置いている")
            # --- 期末 ---
            if m in keep:
                continue
            if m in getsuji:
                yield (tab, "期末棚卸し", m, kimatsu[m], src2,
                       f"{m}末の実地棚卸 {kimatsu[m]:,}。シートR3（=J57+T57）を四捨五入した値。"
                       + besshi)
            else:
                yield (tab, "期末棚卸し", m, kimatsu[m], src2 if tab in GETSUJI else SRC1,
                       f"{m}末に実地棚卸をしていないので、直前の実地 {kimatsu[m]:,} を"
                       "持ち越している（期首＝期末なので売上原価には効かない）")


def check(wb=None):
    import build2
    for tab in KISHU:
        assert "期首棚卸し" in build2.RIDX[tab], f"{tab} に期首棚卸し行が無い"
        assert "期末棚卸し" in build2.RIDX[tab], f"{tab} に期末棚卸し行が無い"
    assert KISHU["業務課"] + KISHU["鳥害対策課"] == 1813683, \
        "ビルメン・鳩の割り振りが元の合計 1,813,683 にならない"
    # ★8月は「期首も期末も入る」か「両方とも入らない」かのどちらかであること。
    #   片方だけだと売上原価が棚卸ぶんまるごとズレる。
    aug = {}
    for tab, plrow, m, _v, _s, _n in rows():
        if m == "8月":
            aug.setdefault(tab, set()).add(plrow)
    for tab, got in aug.items():
        assert got == {"期首棚卸し", "期末棚卸し"}, f"{tab} の8月が片方だけ（{got}）"
    for tab in NO_AUG:
        assert tab not in aug, f"{tab} は8月を入れない約束"
    if wb is None:
        return
    for tab, plrow, m, v, _s, _n in rows():
        c = wb[tab][f"{build2.MCOL[m]}{build2.RIDX[tab][plrow]}"]
        assert not c.value, (f"{tab} {plrow} {m} に既に {c.value} が入っている。"
                             f"棚卸の {v:,} を足すと二重計上になる")


def hold_rows():
    out = [("8月", tab, "★期首棚卸し・期末棚卸しが両方とも空",
            "月次棚卸表がサービスアカウントに共有されておらず404で読めないため、"
            "8月末の実地棚卸が取れない。"
            f"7月末の実地は {GETSUJI[tab]['7月']:,}。"
            "期首だけ入れると 売上原価(a)＝仕入合計＋期首−期末 が棚卸ぶんまるごと多く出るので、"
            "両方とも空のままにしている。"
            "pl-writer@prostyle-pl.iam.gserviceaccount.com に閲覧権限をつければ入ります")
           for tab in NO_AUG]
    out += [("9月〜7月", tab, "期首棚卸し・期末棚卸し",
             f"2025年8月31日の実地棚卸 {KISHU[tab]:,} を9月〜7月へそのまま流している。"
             "月ごとの棚卸表がドライブに見当たらないため。"
             "期首＝期末なので売上原価には効かないが、月次の実地があれば差し替えたい"
             + ("。★8月だけは実地（347,962）が届いたので入れている"
                if tab == "さわら十三里屋" else ""))
            for tab in ("さわら十三里屋", "神栖横丁", "業務課", "鳥害対策課")]
    out += [("9月・10月・11月", "もも焼きJAPAN", "期末棚卸し",
             "既存21期PLの438,111／463,688／484,615をそのまま残している。"
             "毎月ちがう数字なので実際に数えた値とみられるが、ドライブの"
             "棚卸表_もも焼JAPAN.xlsx は20251231からしかシートが無いので裏は取れていない。"
             "利用者確認 2026-09-01「もも焼きの9・10・11月はナシ。既存のままでOK」で"
             "既存PLのままにすると決めた。記録として残す")]
    out += [("9月", "業務課／鳥害対策課", "期首棚卸し",
             "21期_棚卸し表の「ビルメン・鳩」タブ 1,813,683 を、品目を数えて"
             "ビルメン薬剤・資材1,080,271／事務用品15,828／鳩対策の資材345,644／"
             "ネットカーテンセット371,940 に分けた（合計は1円まで一致）。"
             "利用者確認 2026-09-01「ネットカーテンは鳩対策」で"
             "業務課1,096,099／鳥害対策課717,584 とした")]
    return out


# ---------------------------------------------------------------- 22期
# 2026-10-20 利用者「棚卸しも反映して」。
#   期首[9月] = 21期の期末（2026年8月31日）。月次の実地が無い部署は21期の値をそのまま持ち越す。
#   期末[m]   = 各店の棚卸表の m月末シート（R3＝J57+T57 を四捨五入）。無い月は前月を持ち越す。
#   期首[m]   = 期末[m-1]
# 入れるのは「いちばん新しい実地のある月」まで（先の月は空けておく）。
def _kimatsu21():
    out = {}
    for tab, kishu in KISHU.items():
        last = kishu
        for m in MONTHS:
            last = GETSUJI.get(tab, {}).get(m, last)
        out[tab] = last
    return out


KISHU22 = _kimatsu21()   # タコハイ161,096／焼きたて屋267,392／りゅうちゃん387,537／もも焼き227,901／
                         # ハナ182,460／十三里屋347,962／神栖横丁50,357／業務課1,096,099／鳥害対策課717,584

# 9月（2026-09-30 シート。2026-10-20 に読んだ）
#   もも焼き 204359.7／りゅうちゃん 436042.47（ドリンクは別紙 20260930d の小計が本紙7行目 J7 に入っている）
#   タコハイ 248440.23／焼きたて屋 233488.33／ハナ 168223.2
#   十三里屋は「棚卸表_さわら十三里屋_20260831」に 20260930 シートが無い → 持ち越し
GETSUJI22 = {
    "タコとハイボール": {"9月": 248440},
    "焼きたて屋":       {"9月": 233488},
    "りゅうちゃん":     {"9月": 436042},
    "もも焼きJAPAN":   {"9月": 204360},
    "韓国酒場ハナ":     {"9月": 168223},
}
UPTO22 = "9月"   # ここまでの月を入れる（次の月末の棚卸表がそろったら進める）

# 各店の月次棚卸表（Googleドライブの .xlsx）。read22() がサービスアカウントで読む。
# ★十三里屋の「棚卸表_さわら十三里屋_20260831」（1XZ5gGE2Us0DlziyXl6Nq1yHgFmiy6qIpgk8X4pO8oHU）は
#   サービスアカウントに共有されていない（403）。Sheetsコネクタでシート名と R3 を見る。
FILES22 = {
    "焼きたて屋":       "1fJPUbw3WNxtBtEiB8jOlYwqCsdO3-N75",
    "タコとハイボール": "1y2Gc-c3PoLSkqngJT2bZVqWKKWF8UBdI",
    "もも焼きJAPAN":   "1GHGnIXPrzt_KXVS9B4uDn5C0tFkulFII",
    "りゅうちゃん":     "1B4QaxjQ2YUPmfBc9bOn9Jwx3NhrrqQ-J",
    "韓国酒場ハナ":     "16M2TxUwFBvyH-Pye9En34esBs4bUmLko",
}


def _half_up(x):
    import decimal
    return int(decimal.Decimal(str(x)).quantize(decimal.Decimal("1"), rounding=decimal.ROUND_HALF_UP))


def read22(ymd):
    """各店の棚卸表から YYYYMMDD シートの合計を読む → {タブ: (四捨五入した値, R3の生の値) または None}

    ★R3 は =J57+T57。別紙（ドリンク）は本紙の行として入っているので足さない。
      念のため、別紙シート（YYYYMMDDd／YYYYMMDDドリンク…）がある店は、本紙に「別紙」の行があり
      その金額が別紙の J57+T57 と一致することを確かめる。
    """
    import io
    import openpyxl
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload
    from google.oauth2.service_account import Credentials
    creds = Credentials.from_service_account_file(
        os.environ["GOOGLE_SA_KEY_FILE"], scopes=["https://www.googleapis.com/auth/drive.readonly"])
    dr = build("drive", "v3", credentials=creds, cache_discovery=False)
    out = {}
    for tab, fid in FILES22.items():
        buf = io.BytesIO()
        dl = MediaIoBaseDownload(buf, dr.files().get_media(fileId=fid, supportsAllDrives=True))
        done = False
        while not done:
            _, done = dl.next_chunk()
        wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=True)
        if ymd not in wb.sheetnames:
            out[tab] = None
            continue
        ws = wb[ymd]
        r3 = ws["R3"].value
        assert isinstance(r3, (int, float)), (tab, ymd, "R3 が数値でない", r3)
        for name in wb.sheetnames:
            if name != ymd and name.startswith(ymd):
                sub = wb[name]
                v = (sub["J57"].value or 0) + (sub["T57"].value or 0)
                hit = [row for row in ws.iter_rows(min_row=1, max_row=60)
                       if any(isinstance(c.value, str) and "別紙" in c.value for c in row)]
                assert any(abs((c.value or 0) - v) < 0.01 for row in hit for c in row
                           if isinstance(c.value, (int, float))), (tab, name, "別紙の小計が本紙に無い")
        out[tab] = (_half_up(r3), r3)
    return out


def rows22():
    """22期ぶん (タブ, PL行, 月, 金額, 元ファイル, メモ)"""
    months = MONTHS[:MONTHS.index(UPTO22) + 1]
    for tab, kishu in KISHU22.items():
        getsuji = GETSUJI22.get(tab, {})
        src = SRC2.get(tab, SRC1) if tab in GETSUJI22 else SRC1
        last = kishu
        for i, m in enumerate(months):
            if i == 0:
                yield (tab, "期首棚卸し", m, kishu, src,
                       f"21期の期末（2026年8月31日）{kishu:,} を22期の期首に置いている")
            else:
                yield (tab, "期首棚卸し", m, last, src, f"前月（{months[i-1]}）末の棚卸 {last:,}")
            if m in getsuji:
                last = getsuji[m]
                yield (tab, "期末棚卸し", m, last, src,
                       f"{m}末の実地棚卸 {last:,}。シートR3（=J57+T57）を四捨五入した値。" + _BESSHI.get(tab, ""))
            else:
                yield (tab, "期末棚卸し", m, last, src,
                       f"{m}末の実地棚卸が無いので直前の {last:,} を持ち越している（期首＝期末なので売上原価には効かない）")


if __name__ == "__main__":
    import sys
    if "--read" in sys.argv:
        # GOOGLE_SA_KEY_FILE=/tmp/sa.json python3 tanaoroshi.py --read 20261031
        ymd = sys.argv[sys.argv.index("--read") + 1]
        for tab, v in read22(ymd).items():
            print(f"{tab:<10} " + ("（まだシートが無い）" if v is None else f"{v[0]:>10,}   R3={v[1]}"))
        print("★十三里屋は Sheetsコネクタで 1XZ5gGE2Us0DlziyXl6Nq1yHgFmiy6qIpgk8X4pO8oHU のシート名を見る")
        sys.exit()
    import collections
    got = list(rows())
    by = collections.defaultdict(dict)
    for tab, plrow, m, v, _s, _n in got:
        by[tab][(plrow, m)] = v
    print(f"棚卸しセル {len(got)} 件\n")
    for tab in KISHU:
        print(f"■ {tab}")
        for plrow in ("期首棚卸し", "期末棚卸し"):
            print(f"   {plrow}  " + "  ".join(
                f"{m}:{by[tab][(plrow,m)]:>8,}" if (plrow, m) in by[tab] else f"{m}:{'—':>8}"
                for m in MONTHS))
    print(f"\n保留 {len(hold_rows())} 件")
